from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "project-report-th.md"
OUTPUT = ROOT / "docs" / "RepairLink_Project_Report_TH_May_June.docx"
ASSET_DIR = ROOT / "docs" / "report-assets"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 30, 45)
MUTED = RGBColor(95, 104, 116)
HEADER_FILL = "F2F4F7"
BORDER = "D9E2EF"


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Tahoma"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Tahoma")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_para(paragraph, before=0, after=6, line=1.1, align=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_heading(doc: Document, text: str, level: int):
    paragraph = doc.add_paragraph()
    if level == 1:
        style_para(paragraph, before=4, after=12)
        run = paragraph.add_run(text)
        set_run_font(run, size=24, color=INK, bold=True)
    elif level == 2:
        style_para(paragraph, before=16, after=8)
        run = paragraph.add_run(text)
        set_run_font(run, size=16, color=BLUE, bold=True)
    else:
        style_para(paragraph, before=10, after=5)
        run = paragraph.add_run(text)
        set_run_font(run, size=13, color=DARK_BLUE, bold=True)


def add_paragraph(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    style_para(paragraph, after=6)
    for chunk, is_code in split_inline_code(text):
        run = paragraph.add_run(chunk)
        set_run_font(run, size=10.5 if is_code else 11, color=INK)
        if is_code:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")


def add_bullet(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="List Bullet")
    style_para(paragraph, after=4, line=1.167)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=INK)


def add_number(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="List Number")
    style_para(paragraph, after=4, line=1.167)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=INK)


def add_code_block(doc: Document, lines: list[str]):
    if not lines:
        return
    paragraph = doc.add_paragraph()
    style_para(paragraph, before=4, after=8)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(60, 65, 75)


def split_inline_code(text: str):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            yield part[1:-1], True
        elif part:
            yield part, False


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color=BORDER, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.append(grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    col_count = len(rows[0])
    table = doc.add_table(rows=1, cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, HEADER_FILL)
        set_cell_border(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        style_para(paragraph, after=0)
        run = paragraph.add_run(value)
        set_run_font(run, size=10.5, color=DARK_BLUE, bold=True)

    for row_values in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            style_para(paragraph, after=0, line=1.1)
            run = paragraph.add_run(value)
            set_run_font(run, size=9.5, color=INK)

    width = 9360 // col_count
    widths = [width] * col_count
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def set_cell_text(cell, text: str, size=8.5, bold=False, color=INK, align=None):
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    style_para(paragraph, after=0, line=1.0, align=align)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def merge_range(table, row_index: int, start_col: int, end_col: int, text: str, fill: str):
    cell = table.cell(row_index, start_col).merge(table.cell(row_index, end_col))
    set_cell_shading(cell, fill)
    set_cell_border(cell)
    set_cell_margins(cell, 60, 60, 80, 80)
    set_cell_text(cell, text, size=8.5, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_workplan_timeline_table(doc: Document, source_rows: list[list[str]]):
    weeks = [
        "01-May", "02-May", "03-May", "04-May",
        "01-Jun", "02-Jun", "03-Jun", "04-Jun",
    ]

    # Source rows come from the UTF-8 Markdown table, avoiding Thai literals in this script.
    source_tasks = [row[1] for row in source_rows[1:] if len(row) > 1]
    schedule = [
        ("01-May", "01-May", 1, 1),
        ("01-May", "02-May", 1, 2),
        ("02-May", "02-May", 2, 2),
        ("02-May", "03-May", 2, 3),
        ("03-May", "04-May", 3, 4),
        ("04-May", "01-Jun", 4, 5),
        ("01-Jun", "02-Jun", 5, 6),
        ("02-Jun", "02-Jun", 6, 6),
        ("02-Jun", "03-Jun", 6, 7),
        ("03-Jun", "03-Jun", 7, 7),
        ("03-Jun", "04-Jun", 7, 8),
        ("04-Jun", "04-Jun", 8, 8),
    ]
    tasks = [
        (task, *schedule[index])
        for index, task in enumerate(source_tasks[: len(schedule)])
    ]

    paragraph = doc.add_paragraph()
    style_para(paragraph, before=4, after=6)
    run = paragraph.add_run("\u0e15\u0e32\u0e23\u0e32\u0e07\u0e41\u0e1c\u0e19\u0e07\u0e32\u0e19\u0e41\u0e1a\u0e1a Timeline / Gantt (May-Jun)")
    set_run_font(run, size=11, color=DARK_BLUE, bold=True)

    table = doc.add_table(rows=3 + len(tasks), cols=3 + len(weeks))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = [3340, 650, 650] + [590] * len(weeks)
    set_table_geometry(table, widths)

    detail_label = "\u0e23\u0e32\u0e22\u0e25\u0e30\u0e40\u0e2d\u0e35\u0e22\u0e14\u0e07\u0e32\u0e19"
    start_label = "\u0e40\u0e23\u0e34\u0e48\u0e21"
    end_label = "\u0e16\u0e36\u0e07"
    merge_range(table, 0, 0, 0, detail_label, HEADER_FILL)
    merge_range(table, 0, 1, 1, start_label, HEADER_FILL)
    merge_range(table, 0, 2, 2, end_label, HEADER_FILL)
    merge_range(table, 0, 3, 6, "May", "E8F5E9")
    merge_range(table, 0, 7, 10, "June", "E8F5E9")

    for label, start, end in (("May", 3, 6), ("Jun", 7, 10)):
        merge_range(table, 1, start, end, label, "F8FAFC")
    for col in range(3):
        cell = table.cell(1, col)
        set_cell_shading(cell, HEADER_FILL)
        set_cell_border(cell)
        set_cell_text(cell, "", size=8)

    for col, label in enumerate((detail_label, start_label, end_label)):
        cell = table.cell(2, col)
        set_cell_shading(cell, HEADER_FILL)
        set_cell_border(cell)
        set_cell_text(
            cell,
            label,
            size=8.5,
            bold=True,
            color=DARK_BLUE,
            align=WD_ALIGN_PARAGRAPH.CENTER if col else None,
        )

    for week_index, week_label in enumerate(weeks):
        cell = table.cell(2, 3 + week_index)
        set_cell_shading(cell, "FFFFFF")
        set_cell_border(cell)
        set_cell_margins(cell, 50, 50, 30, 30)
        set_cell_text(cell, week_label, size=6.6, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row_index, (task, start_label, end_label, start_week, end_week) in enumerate(tasks, start=3):
        set_cell_text(table.cell(row_index, 0), task, size=7.6)
        set_cell_text(table.cell(row_index, 1), start_label, size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(row_index, 2), end_label, size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
        for col in range(3):
            cell = table.cell(row_index, col)
            set_cell_border(cell)
            set_cell_margins(cell, 45, 45, 60, 60)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        for week_number in range(1, len(weeks) + 1):
            cell = table.cell(row_index, 2 + week_number)
            set_cell_border(cell)
            set_cell_margins(cell, 35, 35, 20, 20)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if start_week <= week_number <= end_week:
                set_cell_shading(cell, "65B741")
            else:
                set_cell_shading(cell, "FFFFFF")
            set_cell_text(cell, "", size=6)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def parse_markdown_table(lines: list[str], start: int):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw_cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r"-+", cell.replace(" ", "")) for cell in raw_cells):
            rows.append(raw_cells)
        index += 1
    return rows, index


def configure_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    for style_name in ("Normal", "List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Tahoma"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
        style._element.rPr.rFonts.set(qn("w:cs"), "Tahoma")
        style.font.size = Pt(11)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("RepairLink Project Report")
    set_run_font(run, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("RepairLink Internal IT Support & Asset Management")
    set_run_font(run, size=9, color=MUTED)


def add_mermaid_image(doc: Document, image_index: int):
    images = [
        ASSET_DIR / "dfd-level-0.png",
        ASSET_DIR / "dfd-level-1.png",
        ASSET_DIR / "use-case-overview.png",
    ]
    if image_index < len(images) and images[image_index].exists():
        doc.add_picture(str(images[image_index]), width=Inches(6.45))
        paragraph = doc.add_paragraph()
        style_para(paragraph, before=4, after=8)
        run = paragraph.add_run(f"Figure {image_index + 1}")
        set_run_font(run, size=9.5, color=MUTED, italic=True)


def build_docx():
    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    doc = Document()
    configure_document(doc)

    in_code = False
    code_lang = ""
    code_lines: list[str] = []
    mermaid_index = 0
    pending_workplan_timeline = False
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_lang = stripped[3:].strip()
                code_lines = []
            else:
                if code_lang == "mermaid":
                    add_mermaid_image(doc, mermaid_index)
                    mermaid_index += 1
                else:
                    add_code_block(doc, code_lines)
                in_code = False
                code_lang = ""
                code_lines = []
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped.startswith("|"):
            table_rows, index = parse_markdown_table(lines, index)
            if pending_workplan_timeline:
                add_workplan_timeline_table(doc, table_rows)
                pending_workplan_timeline = False
            else:
                add_table(doc, table_rows)
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading_match:
            heading_text = heading_match.group(2)
            add_heading(doc, heading_text, len(heading_match.group(1)))
            pending_workplan_timeline = heading_text.startswith("11. Workplan")
            index += 1
            continue

        if stripped.startswith("- "):
            add_bullet(doc, stripped[2:])
            index += 1
            continue

        number_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if number_match:
            add_number(doc, number_match.group(1))
            index += 1
            continue

        add_paragraph(doc, stripped)
        index += 1

    props = doc.core_properties
    props.author = "RepairLink Project Team"
    props.title = "RepairLink Project Report"
    props.subject = "DFD, Use Case, Workplan, Requirement Gathering"
    props.keywords = "RepairLink, DFD, Use Case, Requirement, Workplan"
    props.comments = ""
    doc.save(OUTPUT)

    # Structural smoke check: make sure Thai was not replaced with question marks.
    rebuilt = Document(OUTPUT)
    text = "\n".join(paragraph.text for paragraph in rebuilt.paragraphs[:20])
    if "????????" in text:
        raise RuntimeError("Thai text was replaced with question marks in the generated DOCX.")
    with ZipFile(OUTPUT) as archive:
        archive.testzip()
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
