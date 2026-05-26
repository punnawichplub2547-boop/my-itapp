from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "docs" / "project-report-5sections-th.md"
SOURCE_WORKPLAN_DOCX = ROOT / "docs" / "RepairLink_Project_Report_TH_May_June.docx"
OUT_DOCX = ROOT / "docs" / "RepairLink_Project_Report_TH_5Sections.docx"
ASSET_DIR = ROOT / "docs" / "report-assets"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

COLORS = {
    "navy": "0B2545",
    "blue": "1F4E79",
    "mid_blue": "2E74B5",
    "light_blue": "D9EAF7",
    "pale_blue": "EEF5FB",
    "header_fill": "E8EEF5",
    "border": "B7C9DA",
    "body": "1F2937",
    "muted": "667085",
    "white": "FFFFFF",
    "bar": "2E74B5",
}

THAI_FONT = "TH Sarabun New"
LATIN_FONT = "Calibri"

FIGURES = [
    ASSET_DIR / "dfd-level-0.png",
    ASSET_DIR / "dfd-level-1.png",
    ASSET_DIR / "use-case-overview.png",
]

WEEK_BUCKETS = [
    ("01-May", "1-7 May", datetime(2026, 5, 1), datetime(2026, 5, 7)),
    ("02-May", "8-14 May", datetime(2026, 5, 8), datetime(2026, 5, 14)),
    ("03-May", "15-21 May", datetime(2026, 5, 15), datetime(2026, 5, 21)),
    ("04-May", "22-31 May", datetime(2026, 5, 22), datetime(2026, 5, 31)),
    ("01-Jun", "1-7 Jun", datetime(2026, 6, 1), datetime(2026, 6, 7)),
    ("02-Jun", "8-14 Jun", datetime(2026, 6, 8), datetime(2026, 6, 14)),
    ("03-Jun", "15-21 Jun", datetime(2026, 6, 15), datetime(2026, 6, 21)),
    ("04-Jun", "22-30 Jun", datetime(2026, 6, 22), datetime(2026, 6, 30)),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = COLORS["border"], size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = mar.find(qn("w:" + key))
        if node is None:
            node = OxmlElement("w:" + key)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_font(run, size: int | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = THAI_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), THAI_FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, size=16, color=COLORS["body"], bold=False) -> None:
    for run in paragraph.runs:
        set_font(run, size=size, bold=bold, color=color)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1


def add_paragraph(doc: Document, text: str = "", size=16, color=COLORS["body"], bold=False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    return paragraph


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = THAI_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), THAI_FONT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    normal.font.size = Pt(16)
    normal.font.color.rgb = RGBColor.from_string(COLORS["body"])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    style_specs = {
        "Title": (26, COLORS["navy"], 0, 8),
        "Subtitle": (16, COLORS["muted"], 0, 12),
        "Heading 1": (22, COLORS["navy"], 16, 8),
        "Heading 2": (19, COLORS["blue"], 12, 6),
        "Heading 3": (17, COLORS["blue"], 8, 4),
    }
    for name, (size, color, before, after) in style_specs.items():
        style = styles[name]
        style.font.name = THAI_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), THAI_FONT)
        style._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_p.add_run("RepairLink Project Report")
    set_font(run, size=10, color=COLORS["muted"])

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("RepairLink")
    set_font(run, size=10, color=COLORS["muted"])


def add_cover(doc: Document, title: str, subtitle: str, lead: str) -> None:
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_font(run, size=30, bold=True, color=COLORS["navy"])
    p.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    set_font(run, size=18, color=COLORS["blue"])
    p.paragraph_format.space_after = Pt(18)

    info = doc.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(info, [2300, 6000], indent_dxa=520)
    rows = [
        ("Project", "RepairLink - Internal IT Support and Asset Management"),
        ("Prepared for", "IT Administrator / IT Manager"),
        ("Prepared date", "25 May 2026"),
        ("Source", "Project report adapted to five-section academic structure"),
    ]
    for row, (left, right) in zip(info.rows, rows):
        row.cells[0].text = left
        row.cells[1].text = right
        set_cell_shading(row.cells[0], COLORS["header_fill"])
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                style_paragraph(paragraph, size=13)
            row.cells[0].paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(callout, [8500], indent_dxa=430)
    cell = callout.cell(0, 0)
    set_cell_shading(cell, COLORS["pale_blue"])
    cell.text = lead
    for paragraph in cell.paragraphs:
        style_paragraph(paragraph, size=15, color=COLORS["navy"])
    doc.add_page_break()


def read_workplan_rows() -> list[tuple[str, str, str]]:
    doc = Document(SOURCE_WORKPLAN_DOCX)
    if len(doc.tables) < 2:
        raise RuntimeError("Expected workplan table in source report.")
    table = doc.tables[1]
    rows: list[tuple[str, str, str]] = []
    for row in table.rows[3:]:
        values = [cell.text.strip() for cell in row.cells[:3]]
        if len(values) == 3 and values[0] and values[1] and values[2]:
            rows.append((values[0], values[1], values[2]))
    if not rows:
        raise RuntimeError("No workplan rows found.")
    return rows


def parse_work_date(value: str) -> datetime:
    cleaned = value.strip()
    for fmt in ("%d-%b", "%d-%B"):
        try:
            parsed = datetime.strptime(cleaned, fmt)
            return parsed.replace(year=2026)
        except ValueError:
            pass
    raise RuntimeError("Could not parse workplan date: {}".format(value))


def format_range(start: str, end: str) -> str:
    if start == end:
        return start
    return "{} - {}".format(start, end)


def add_workplan(doc: Document, rows: list[tuple[str, str, str]]) -> None:
    add_paragraph(doc, "Workplan: May - June 2026", size=17, bold=True, color=COLORS["navy"])
    table = doc.add_table(rows=3 + len(rows), cols=11)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [3050, 860, 860] + [570] * 8
    set_table_geometry(table, widths)

    # Header row 1
    headers = ["Task", "Start", "End"] + ["Q2"] * 8
    for idx, value in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = value
        set_cell_shading(cell, COLORS["navy"] if idx < 3 else COLORS["blue"])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_paragraph(p, size=11 if idx >= 3 else 12, color=COLORS["white"], bold=True)

    # Header row 2
    row = table.rows[1]
    for idx, value in enumerate(["", "", "", "May", "May", "May", "May", "Jun", "Jun", "Jun", "Jun"]):
        cell = row.cells[idx]
        cell.text = value
        set_cell_shading(cell, COLORS["header_fill"])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_paragraph(p, size=11, color=COLORS["navy"], bold=True)

    # Header row 3
    row = table.rows[2]
    for idx, value in enumerate(["", "", ""] + ["{}\n{}".format(code, label) for code, label, _, _ in WEEK_BUCKETS]):
        cell = row.cells[idx]
        cell.text = value
        set_cell_shading(cell, COLORS["pale_blue"] if idx >= 3 else COLORS["header_fill"])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_paragraph(p, size=9 if idx >= 3 else 10, color=COLORS["navy"], bold=True)

    for row_idx, (task, start_text, end_text) in enumerate(rows, start=3):
        row = table.rows[row_idx]
        base = [task, start_text, end_text]
        start = parse_work_date(start_text)
        end = parse_work_date(end_text)
        for col_idx, text in enumerate(base):
            cell = row.cells[col_idx]
            cell.text = text
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                style_paragraph(p, size=10 if col_idx == 0 else 9, color=COLORS["body"])
        for week_idx, (_, _, week_start, week_end) in enumerate(WEEK_BUCKETS, start=3):
            cell = row.cells[week_idx]
            overlaps = start <= week_end and end >= week_start
            cell.text = format_range(start_text, end_text) if overlaps else ""
            if overlaps:
                set_cell_shading(cell, COLORS["bar"])
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                style_paragraph(p, size=7, color=COLORS["white"] if overlaps else COLORS["body"], bold=overlaps)

    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell, top=80, bottom=80, start=90, end=90)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    add_paragraph(doc, "The blue cells identify the actual work period recorded in the latest project report file.", size=12, color=COLORS["muted"])


def add_markdown_table(doc: Document, table_lines: list[str]) -> None:
    rows = []
    for line in table_lines:
        parts = [part.strip() for part in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", part or "---") for part in parts):
            continue
        rows.append(parts)
    if not rows:
        return
    max_cols = max(len(row) for row in rows)
    for row in rows:
        row.extend([""] * (max_cols - len(row)))
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if max_cols == 2:
        widths = [2800, PAGE_WIDTH_DXA - 2800]
    elif max_cols == 3:
        widths = [2100, 3300, PAGE_WIDTH_DXA - 5400]
    else:
        widths = [PAGE_WIDTH_DXA // max_cols] * max_cols
        widths[-1] = PAGE_WIDTH_DXA - sum(widths[:-1])
    set_table_geometry(table, widths)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            set_cell_shading(cell, COLORS["header_fill"] if r_idx == 0 else COLORS["white"])
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                style_paragraph(paragraph, size=13, bold=r_idx == 0, color=COLORS["navy"] if r_idx == 0 else COLORS["body"])
    doc.add_paragraph()


def add_figure(doc: Document, image_path: Path, idx: int) -> None:
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(5.9))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run("Figure {}".format(idx))
    set_font(run, size=12, color=COLORS["muted"])
    caption.paragraph_format.space_after = Pt(10)


def split_intro(lines: list[str]) -> tuple[str, str, str, list[str]]:
    title = ""
    subtitle = ""
    intro: list[str] = []
    rest: list[str] = []
    in_intro = True
    for line in lines:
        if line.startswith("## "):
            in_intro = False
        if in_intro:
            if line.startswith("# "):
                title = line[2:].strip()
            elif line.strip() and not subtitle:
                subtitle = line.strip()
            elif line.strip():
                intro.append(line.strip())
        else:
            rest.append(line)
    lead = " ".join(intro).strip()
    return title, subtitle, lead, rest


def build_document() -> None:
    markdown = SOURCE_MD.read_text(encoding="utf-8")
    if "????????" in markdown:
        raise RuntimeError("Question-mark replacement detected in markdown source.")
    lines = markdown.splitlines()
    title, subtitle, lead, body_lines = split_intro(lines)
    workplan_rows = read_workplan_rows()

    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)
    add_cover(doc, title, subtitle, lead)

    table_buffer: list[str] = []
    in_mermaid = False
    figure_idx = 0

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            add_markdown_table(doc, table_buffer)
            table_buffer = []

    for raw_line in body_lines:
        line = raw_line.rstrip()
        if line.startswith("```mermaid"):
            flush_table()
            in_mermaid = True
            figure_idx += 1
            add_figure(doc, FIGURES[figure_idx - 1], figure_idx)
            continue
        if in_mermaid:
            if line.startswith("```"):
                in_mermaid = False
            continue
        if not line.strip():
            flush_table()
            continue
        if line.startswith("|"):
            table_buffer.append(line)
            continue
        flush_table()

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            heading = line[4:].strip()
            doc.add_heading(heading, level=2)
            if "Workplan" in heading:
                add_workplan(doc, workplan_rows)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(line[2:].strip())
            set_font(run, size=16, color=COLORS["body"])
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.1
        else:
            add_paragraph(doc, line.strip())

    flush_table()

    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            for run in paragraph.runs:
                set_font(run, bold=True)

    if "????????" in "\n".join(p.text for p in doc.paragraphs):
        raise RuntimeError("Question-mark replacement detected in document paragraphs.")

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)


if __name__ == "__main__":
    build_document()
    print(OUT_DOCX)
