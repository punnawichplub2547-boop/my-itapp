from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "docs" / "RepairLink_Project_Report_TH_May_June.docx"

NAVY = "1F4E79"
DARK_NAVY = "203864"
MID_BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
GRID = "9EADBD"
WHITE = "FFFFFF"
TEXT = RGBColor(31, 43, 55)
MUTED = RGBColor(80, 95, 115)
HEADING = RGBColor(31, 78, 121)


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Tahoma"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Tahoma")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_text(cell, text, size=8.0, color=TEXT, bold=False, align=None):
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def set_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    for existing in list(tc_pr.findall(qn("w:shd"))):
        tc_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_border(cell, color=GRID, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=60, bottom=60, start=80, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def parse_day_month(value: str) -> date | None:
    text = value.strip()
    for fmt in ("%d-%b", "%d-%B"):
        try:
            parsed = datetime.strptime(text, fmt)
            return date(2026, parsed.month, parsed.day)
        except ValueError:
            continue
    return None


def overlaps(left_start: date, left_end: date, right_start: date, right_end: date) -> bool:
    return left_start <= right_end and right_start <= left_end


def polish_general_styles(document: Document):
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if text.startswith("#"):
            continue
        if text[:2].isdigit() and ". " in text[:5]:
            for run in paragraph.runs:
                set_run_font(run, color=HEADING, bold=True)
        elif text.startswith("Figure"):
            for run in paragraph.runs:
                set_run_font(run, size=9.0, color=MUTED, italic=True)

    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                set_border(cell)
                set_cell_margins(cell)
                if row_index == 0:
                    set_shading(cell, NAVY)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            set_run_font(run, color=RGBColor(255, 255, 255), bold=True)


def polish_workplan_table(table):
    week_ranges = [
        (date(2026, 5, 1), date(2026, 5, 7)),
        (date(2026, 5, 8), date(2026, 5, 14)),
        (date(2026, 5, 15), date(2026, 5, 21)),
        (date(2026, 5, 22), date(2026, 5, 31)),
        (date(2026, 6, 1), date(2026, 6, 7)),
        (date(2026, 6, 8), date(2026, 6, 14)),
        (date(2026, 6, 15), date(2026, 6, 21)),
        (date(2026, 6, 22), date(2026, 6, 30)),
    ]
    week_labels = [
        "01-May\n1-7 May",
        "02-May\n8-14 May",
        "03-May\n15-21 May",
        "04-May\n22-31 May",
        "01-Jun\n1-7 Jun",
        "02-Jun\n8-14 Jun",
        "03-Jun\n15-21 Jun",
        "04-Jun\n22-30 Jun",
    ]

    # Top quarter/month rows.
    for col in range(len(table.columns)):
        for row_index in (0, 1):
            cell = table.cell(row_index, col)
            set_border(cell, DARK_NAVY)
            set_cell_margins(cell, 50, 50, 40, 40)
            if col < 3:
                set_shading(cell, DARK_NAVY)
            else:
                set_shading(cell, NAVY if row_index == 0 else PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=8.5,
                        color=RGBColor(255, 255, 255) if row_index == 0 or col < 3 else HEADING,
                        bold=True,
                    )

    # Column header row with clearer date ranges.
    for col in range(3):
        cell = table.cell(2, col)
        set_shading(cell, NAVY)
        set_border(cell, DARK_NAVY)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, color=RGBColor(255, 255, 255), bold=True, size=8.2)

    for offset, label in enumerate(week_labels):
        cell = table.cell(2, 3 + offset)
        set_shading(cell, LIGHT_BLUE)
        set_border(cell, GRID)
        set_cell_margins(cell, 40, 40, 20, 20)
        set_paragraph_text(cell, label, size=6.2, bold=True, color=HEADING, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Task rows: rebuild the bar based on user-edited dates and put exact date text on the bar.
    for row_index in range(3, len(table.rows)):
        row = table.rows[row_index]
        task = row.cells[0].text.strip()
        start_text = row.cells[1].text.strip()
        end_text = row.cells[2].text.strip()
        start = parse_day_month(start_text)
        end = parse_day_month(end_text)

        for col in range(3):
            set_shading(row.cells[col], "FFFFFF")
            set_border(row.cells[col], GRID)
            set_cell_margins(row.cells[col], 45, 45, 60, 60)
        set_paragraph_text(row.cells[0], task, size=7.4, color=TEXT)
        set_paragraph_text(row.cells[1], start_text, size=7.2, color=TEXT, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_paragraph_text(row.cells[2], end_text, size=7.2, color=TEXT, align=WD_ALIGN_PARAGRAPH.CENTER)

        for col in range(3, 11):
            set_shading(row.cells[col], "FFFFFF")
            set_border(row.cells[col], GRID)
            set_cell_margins(row.cells[col], 35, 35, 20, 20)
            set_paragraph_text(row.cells[col], "", size=6)

        if start is None or end is None:
            continue
        if end < start:
            start, end = end, start

        active_cols = [
            3 + offset
            for offset, (week_start, week_end) in enumerate(week_ranges)
            if overlaps(start, end, week_start, week_end)
        ]
        if not active_cols:
            continue

        start_col, end_col = active_cols[0], active_cols[-1]
        bar_cell = row.cells[start_col]
        if end_col > start_col:
            bar_cell = bar_cell.merge(row.cells[end_col])
        set_shading(bar_cell, MID_BLUE)
        set_border(bar_cell, DARK_NAVY)
        set_cell_margins(bar_cell, 35, 35, 40, 40)
        range_text = start_text if start_text == end_text else f"{start_text} - {end_text}"
        set_paragraph_text(
            bar_cell,
            range_text,
            size=6.8,
            bold=True,
            color=RGBColor(255, 255, 255),
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )


def main():
    document = Document(DOCX_PATH)
    polish_general_styles(document)
    polish_workplan_table(document.tables[1])
    document.save(DOCX_PATH)

    # Structural checks.
    rebuilt = Document(DOCX_PATH)
    texts = [paragraph.text for paragraph in rebuilt.paragraphs]
    for table in rebuilt.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    joined = "\n".join(texts)
    if "????????" in joined:
        raise RuntimeError("Question-mark replacement detected after styling.")
    with ZipFile(DOCX_PATH) as archive:
        archive.testzip()
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
